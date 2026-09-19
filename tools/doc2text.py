#!/usr/bin/env python3
"""doc2text: local reference library for token-efficient document lookups.

Download web pages / documents ONCE, convert them to page-marked plain text,
flag pages where text extraction is unreliable, and keep an index so a reader
can grep the text and open only the lines it needs.

Usage (paths are relative to the project root; run from anywhere):
  python tools/doc2text.py fetch URL [--name ID] [--note TEXT] [--force]
  python tools/doc2text.py add PATH  [--name ID] [--note TEXT] [--source URL]
  python tools/doc2text.py render ID PAGES [--dpi N]   PAGES like 44 or 40-44,46 (PDF page numbers)
  python tools/doc2text.py note ID TEXT
  python tools/doc2text.py rebuild                     reconvert every document
  python tools/doc2text.py list

Layout:
  reference/raw/         downloaded originals (local files added with 'add' stay where they are)
  reference/text/ID.txt  converted text with '=== PAGE n | label | flags ===' markers
  reference/render/      PNG renders of single pages, made on demand
  reference/manifest.json, reference/INDEX.md (generated)

Page flags:
  math    equation symbols recovered via Unicode NFKC, but sub/superscripts are flattened
  garbled font encoding is broken; text is junk -> render the page
  lowtext little or no text (scan, figure page, or JavaScript-rendered web page)
  img     page contains embedded images (figures/graphs are not in the text)
"""
import argparse
import datetime
import hashlib
import json
import re
import struct
import sys
import unicodedata
import zlib
from pathlib import Path
from urllib.parse import urljoin, urlparse

ROOT = Path(__file__).resolve().parent.parent
REF = ROOT / "reference"
RAW, TEXT, RENDER = REF / "raw", REF / "text", REF / "render"
MANIFEST, INDEX = REF / "manifest.json", REF / "INDEX.md"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
DOC_LINK = re.compile(r"\.(pdf|docx?|pptx?|xlsx?|csv|zip)(\?|$)|DownloadDocument", re.I)
MATH_CHARS = re.compile("[" + chr(0x1D400) + "-" + chr(0x1D7FF) + chr(0x210E) + "]")  # math alphanumerics U+1D400-1D7FF and U+210E (italic h)


def now():
    return datetime.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M UTC%z")


def slug(s):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("_.")[:80] or "doc"


def rel(p):
    p = Path(p).resolve()
    try:
        return p.relative_to(ROOT).as_posix()
    except ValueError:
        return str(p)


def load_manifest():
    return json.loads(MANIFEST.read_text(encoding="utf-8")) if MANIFEST.exists() else {}


def save_manifest(m):
    REF.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(m, indent=2, sort_keys=True), encoding="utf-8")
    write_index(m)


def sniff(data):
    """Identify the file type from its bytes (works on truncated Office files too)."""
    if data[:4] == b"%PDF":
        return "pdf"
    if data[:4] == b"\xd0\xcf\x11\xe0":
        return "ole"  # password-encrypted Office file or legacy .doc/.ppt/.xls
    if data[:2] == b"PK":
        for marker, kind in ((b"word/document.xml", "docx"), (b"ppt/slides/", "pptx"), (b"xl/workbook", "xlsx")):
            if marker in data:
                return kind
        return "zip"
    head = data[:4000].lower()
    if b"<html" in head or b"<!doctype html" in head:
        return "html"
    return "txt"


# ---------------------------------------------------------------- PDF

def pdf_layout_text(page):
    """Rebuild reading-order lines from word boxes; keep wide gaps so tables stay aligned."""
    words = page.get_text("words")  # (x0, y0, x1, y1, text, block, line, word)
    if not words:
        return ""
    widths = sorted((w[2] - w[0]) / max(len(w[4]), 1) for w in words)
    cw = widths[len(widths) // 2] or 5.0  # typical character width in points
    heights = sorted(w[3] - w[1] for w in words)
    hm = heights[len(heights) // 2] or 10.0  # typical line height; math fonts have oversized boxes
    x_min = min(w[0] for w in words)
    rows = []  # [y_center, [words]]
    for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        yc = (w[1] + w[3]) / 2
        if rows and abs(rows[-1][0] - yc) < max(2.0, 0.45 * hm):
            rows[-1][1].append(w)
        else:
            rows.append([yc, [w]])
    lines, prev_y = [], None
    for yc, ws in rows:
        if prev_y is not None and yc - prev_y > 1.8 * hm:
            lines.append("")
        prev_y = yc
        s, last_x1 = "", None
        for w in sorted(ws, key=lambda w: w[0]):
            col = int((w[0] - x_min) / cw)
            if last_x1 is None:
                s = " " * col + w[4]
            elif w[0] - last_x1 > 2.5 * cw:  # real column gap -> pad to position
                s += " " * max(col - len(s), 2) + w[4]
            else:
                s += " " + w[4]
            last_x1 = w[2]
        lines.append(s.rstrip())
    return "\n".join(lines)


def garbled(text):
    toks = text.split()
    if len(toks) < 20:
        return False
    ok = 0
    for t in toks:
        core = t.strip(".,;:()[]{}\"'!?*")
        if re.fullmatch(r"[A-Za-z][A-Za-z'/-]+|[AaI]|\d[\d.,:/-]*%?|[<>=+-]", core):
            ok += 1
    return ok / len(toks) < 0.5


def printed_page(text):
    tail = [l for l in text.splitlines() if l.strip()][-3:]
    for line in tail:
        m = re.search(r"Page:?\s*(\d+)\s*$", line)
        if m:
            return m.group(1)
    return ""


def has_figure(page):
    """True if the page shows an image big enough to be a figure (ignores icons and image-drawn numbers)."""
    for img in page.get_images():
        for r in page.get_image_rects(img[0]):
            if r.width > 60 and r.height > 40:
                return True
    return False


def pdf_pages(path):
    import pymupdf
    doc = pymupdf.open(str(path))
    pages = []
    for i, page in enumerate(doc, 1):
        raw = pdf_layout_text(page)
        text = unicodedata.normalize("NFKC", raw)
        flags = []
        if MATH_CHARS.search(raw):
            flags.append("math")
        if len(text.strip()) < 100:
            flags.append("lowtext")
        elif garbled(text):
            flags.append("garbled")
        if has_figure(page):
            flags.append("img")
        label = page.get_label() or printed_page(text)
        pages.append((i, f"p.{label}" if label else "", flags, text))
    toc = [f"{'  ' * (lvl - 1)}{title} -> PDF page {pno}" for lvl, title, pno in doc.get_toc()]
    return pages, toc


# ---------------------------------------------------------------- HTML

def html_pages(data, source=""):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, "lxml")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    links = []
    for a in soup.find_all("a", href=True):
        href = urljoin(source, a["href"]) if source else a["href"]
        if DOC_LINK.search(href) and href not in links:
            links.append(href)
    # not <form>: ASP.NET sites (e.g. saeaerodesign.com) wrap the whole page in one
    for t in soup(["script", "style", "noscript", "nav", "header", "footer", "aside", "svg", "iframe"]):
        t.decompose()
    for sel in ("[role=navigation]", ".navbox", ".mw-editsection", "#mw-navigation", "sup.reference", ".reflist"):
        for t in soup.select(sel):
            t.decompose()
    main = (soup.find("main") or soup.find("article") or soup.find(attrs={"role": "main"})
            or soup.find(id="mw-content-text") or soup.body or soup)
    for tr in main.find_all("tr"):
        cells = [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])]
        tr.replace_with(soup.new_string("\n| " + " | ".join(cells) + " |\n"))
    for lvl in range(1, 7):
        for h in main.find_all(f"h{lvl}"):
            h.replace_with(soup.new_string("\n\n" + "#" * lvl + " " + h.get_text(" ", strip=True) + "\n"))
    for li in main.find_all("li"):
        li.insert(0, soup.new_string("\n- "))
    for br in main.find_all("br"):
        br.replace_with(soup.new_string("\n"))
    for b in main.find_all(["p", "div", "section", "blockquote", "pre", "dd", "dt", "table"]):
        b.append(soup.new_string("\n"))
    text = unicodedata.normalize("NFKC", main.get_text(""))
    text = re.sub(r"[ \t\r\f\v\xa0]+", " ", text)
    text = "\n".join(l.strip() for l in text.splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if title:
        text = f"# {title}\n\n{text}"
    if links:
        text += "\n\n## Document links on this page\n" + "\n".join(f"- {u}" for u in links[:100])
    flags = ["lowtext"] if len(text) < 300 else []
    return [(1, "", flags, text)]


# ---------------------------------------------------------------- Office

def docx_pages(path):
    import docx
    d = docx.Document(str(path))
    items = d.iter_inner_content() if hasattr(d, "iter_inner_content") else list(d.paragraphs) + list(d.tables)
    out = []
    for it in items:
        if hasattr(it, "rows"):
            for r in it.rows:
                out.append("| " + " | ".join(c.text.strip() for c in r.cells) + " |")
        elif it.text.strip():
            style = (it.style.name if it.style is not None else "").lower()
            out.append(("## " if style.startswith("heading") else "") + it.text.strip())
    return [(1, "", [], "\n".join(out))]


def pptx_shape_text(shapes, out):
    for sh in shapes:
        if hasattr(sh, "shapes"):  # group shape
            pptx_shape_text(sh.shapes, out)
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            out += [l.strip() for l in sh.text_frame.text.splitlines() if l.strip()]
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                out.append("| " + " | ".join(c.text.strip() for c in r.cells) + " |")


def pptx_pages(path):
    try:
        from pptx import Presentation
        pages = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            out = []
            pptx_shape_text(slide.shapes, out)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    out.append("[speaker notes] " + notes)
            pages.append((i, "", [], "\n".join(out)))
        return pages
    except Exception:
        return pptx_raw_scan(path)


def pptx_raw_scan(path):
    """Recover slide text from a truncated/corrupt pptx by reading zip local headers directly."""
    d, i, found = Path(path).read_bytes(), 0, {}
    while (i := d.find(b"PK\x03\x04", i)) >= 0:
        try:
            _, _, meth, _, _, _, csz, _, nl, el = struct.unpack("<HHHHHIIIHH", d[i + 4:i + 30])
            name = d[i + 30:i + 30 + nl].decode("utf-8", "replace")
            m = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
            if m:
                start = i + 30 + nl + el
                blob = d[start:start + csz] if csz else d[start:start + 4_000_000]
                xml = zlib.decompressobj(-15).decompress(blob) if meth == 8 else blob
                paras = re.findall(r"<a:p>(.*?)</a:p>", xml.decode("utf-8", "replace"), re.S)
                lines = ["".join(re.findall(r"<a:t>([^<]*)</a:t>", p)) for p in paras]
                found[int(m.group(1))] = "\n".join(l for l in lines if l.strip())
        except Exception:
            pass
        i += 4
    return [(n, "", ["truncated-file"], found[n]) for n in sorted(found)]


def xlsx_pages(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets, 1):
        rows = ["\t".join("" if v is None else str(v) for v in r).rstrip()
                for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
        pages.append((i, ws.title, [], "\n".join(rows)))
    return pages


# ---------------------------------------------------------------- conversion + index

def convert(doc_id, e):
    path = Path(e["path"]) if Path(e["path"]).is_absolute() else ROOT / e["path"]
    data = path.read_bytes()
    kind, toc = sniff(data), []
    if kind == "pdf":
        pages, toc = pdf_pages(path)
    elif kind == "html":
        pages = html_pages(data, e.get("source", ""))
    elif kind == "docx":
        pages = docx_pages(path)
    elif kind == "pptx":
        pages = pptx_pages(path)
    elif kind == "xlsx":
        pages = xlsx_pages(path)
    elif kind == "ole":
        pages = [(1, "", ["encrypted-or-legacy"], "(not converted: password-encrypted or legacy binary Office file)")]
    elif kind == "zip":
        pages = [(1, "", ["not-converted"], "(zip archive: extract it and add the files individually)")]
    else:
        text = data.decode("utf-8", "replace")
        pages = [(1, "", ["lowtext"] if len(text) < 100 else [], text)]

    flags = {}
    for n, _, fl, _ in pages:
        for f in fl:
            flags.setdefault(f, []).append(n)
    chars = sum(len(t) for *_, t in pages)
    e.update(type=kind, pages=len(pages), chars=chars, tokens=chars // 4, flags=flags,
             sha256=hashlib.sha256(data).hexdigest()[:16], converted=now())

    head = [f"# ID: {doc_id}", f"# SOURCE: {e.get('source') or '(local file)'}", f"# ORIGINAL: {e['path']}",
            f"# CONVERTED: {e['converted']} | type {kind} | {len(pages)} pages | ~{chars // 4} tokens"]
    if e.get("note"):
        head.append(f"# NOTE: {e['note']}")
    for f, pn in flags.items():
        head.append(f"# FLAG {f}: pages {compact(pn)}")
    if toc:
        head += ["# TOC:"] + [f"#   {t}" for t in toc]
    body = []
    for n, label, fl, text in pages:
        parts = [f"PAGE {n}"] + ([label] if label else []) + ([", ".join(fl)] if fl else [])
        body += ["", "=== " + " | ".join(parts) + " ===", text.rstrip()]
    TEXT.mkdir(parents=True, exist_ok=True)
    (TEXT / f"{doc_id}.txt").write_text("\n".join(head + body) + "\n", encoding="utf-8")


def compact(nums):
    """[1,2,3,7,9,10] -> '1-3,7,9-10'"""
    out, nums = [], sorted(set(nums))
    for n in nums:
        if out and n == out[-1][1] + 1:
            out[-1][1] = n
        else:
            out.append([n, n])
    return ",".join(str(a) if a == b else f"{a}-{b}" for a, b in out)


def parse_pages(spec):
    pages = []
    for part in spec.split(","):
        a, _, b = part.partition("-")
        pages += range(int(a), int(b or a) + 1)
    return pages


def write_index(m):
    lines = [
        "# Reference library index",
        "",
        "Generated by `tools/doc2text.py`; do not edit by hand. Add notes with `python tools/doc2text.py note ID TEXT`.",
        "",
        "How to use (token-efficient):",
        "1. Read this index, not whole documents.",
        "2. Grep `reference/text/` for keywords, then Read only the matching line ranges.",
        "3. Page markers look like `=== PAGE 44 | p.36 | math ===` (PDF page | printed page | flags).",
        "   Grep `=== PAGE 44 ` to jump to a page; cite the printed page.",
        "4. Flags: math = equation recovered but sub/superscripts flattened; garbled = broken font, render it;",
        "   lowtext = little text (scan/figure/JS page); img = figures not in text; truncated-file = partial recovery.",
        "5. Render a page image only when a flag or a figure requires it: `python tools/doc2text.py render ID PAGE`.",
        "",
        "| ID | Type | Pages | ~Tokens | Flags | Source | Note |",
        "|---|---|---|---|---|---|---|",
    ]
    for doc_id in sorted(m):
        e = m[doc_id]
        fl = "; ".join(f"{f} {compact(p)}" for f, p in e.get("flags", {}).items())
        src = e.get("source") or e["path"]
        note = (e.get("note") or "").replace("|", "/")
        lines.append(f"| {doc_id} | {e.get('type', '?')} | {e.get('pages', '?')} | {e.get('tokens', '?')} | {fl} | {src} | {note} |")
    INDEX.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- commands

def cmd_fetch(a, m):
    import warnings
    warnings.filterwarnings("ignore", message=".*doesn't match a supported version")
    import requests
    doc_id = slug(a.name or (Path(urlparse(a.url).path).stem + ("_" + urlparse(a.url).query if urlparse(a.url).query else "")))
    existing = m.get(doc_id)
    if existing and (ROOT / existing["path"]).exists() and not a.force:
        print(f"[OK] {doc_id}: already downloaded ({existing['path']}); use --force to re-download")
        return
    try:
        r = requests.get(a.url, headers=UA, timeout=60)
        r.raise_for_status()
        content = r.content
    except requests.exceptions.SSLError:
        content = curl_get(a.url)  # some university servers send incomplete cert chains; curl copes
    except requests.RequestException as ex:
        sys.exit(f"[X] {doc_id}: fetch failed: {type(ex).__name__}: {str(ex)[:200]}")
    kind = sniff(content)
    ext = {"ole": "bin", "txt": Path(urlparse(a.url).path).suffix.lstrip(".") or "txt"}.get(kind, kind)
    RAW.mkdir(parents=True, exist_ok=True)
    dest = RAW / f"{doc_id}.{ext}"
    dest.write_bytes(content)
    e = m.setdefault(doc_id, {})
    e.update(path=rel(dest), source=a.url, fetched=now())
    if a.note:
        e["note"] = a.note
    convert(doc_id, e)
    report(doc_id, e)


def curl_get(url):
    import shutil
    import subprocess
    curl = shutil.which("curl") or sys.exit(f"[X] SSL error and curl not found: {url}")
    p = subprocess.run([curl, "-sfL", "-A", UA["User-Agent"], url], capture_output=True, timeout=300)
    if p.returncode != 0 or not p.stdout:
        sys.exit(f"[X] fetch failed (SSL error, curl fallback exit {p.returncode}): {url}")
    return p.stdout


def cmd_add(a, m):
    path = Path(a.path).resolve()
    if not path.exists():
        sys.exit(f"[X] not found: {a.path}")
    doc_id = slug(a.name or path.stem)
    e = m.setdefault(doc_id, {})
    e.update(path=rel(path), added=now())
    if a.source:
        e["source"] = a.source
    if a.note:
        e["note"] = a.note
    convert(doc_id, e)
    report(doc_id, e)


def cmd_render(a, m):
    import pymupdf
    e = m.get(a.id) or sys.exit(f"[X] unknown ID: {a.id}")
    doc = pymupdf.open(str(ROOT / e["path"]))
    RENDER.mkdir(parents=True, exist_ok=True)
    for n in parse_pages(a.pages):
        out = RENDER / f"{a.id}_p{n}.png"
        doc[n - 1].get_pixmap(dpi=a.dpi).save(str(out))
        print(f"[OK] {rel(out)}")


def cmd_note(a, m):
    e = m.get(a.id) or sys.exit(f"[X] unknown ID: {a.id}")
    e["note"] = a.text
    convert(a.id, e)
    print(f"[OK] note set on {a.id}")


def cmd_rebuild(a, m):
    for doc_id, e in m.items():
        convert(doc_id, e)
        report(doc_id, e)


def cmd_list(a, m):
    print(INDEX.read_text(encoding="utf-8") if INDEX.exists() else "(empty library)")


def report(doc_id, e):
    fl = "; ".join(f"{f} {compact(p)}" for f, p in e["flags"].items()) or "none"
    print(f"[OK] {doc_id}: {e['type']}, {e['pages']} pages, ~{e['tokens']} tokens, flags: {fl} -> reference/text/{doc_id}.txt")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    f = sub.add_parser("fetch"); f.add_argument("url"); f.add_argument("--name"); f.add_argument("--note"); f.add_argument("--force", action="store_true")
    ad = sub.add_parser("add"); ad.add_argument("path"); ad.add_argument("--name"); ad.add_argument("--note"); ad.add_argument("--source")
    r = sub.add_parser("render"); r.add_argument("id"); r.add_argument("pages"); r.add_argument("--dpi", type=int, default=110)
    n = sub.add_parser("note"); n.add_argument("id"); n.add_argument("text")
    sub.add_parser("rebuild"); sub.add_parser("list")
    a = p.parse_args()
    m = load_manifest()
    {"fetch": cmd_fetch, "add": cmd_add, "render": cmd_render, "note": cmd_note,
     "rebuild": cmd_rebuild, "list": cmd_list}[a.cmd](a, m)
    if a.cmd != "list":
        save_manifest(m)


if __name__ == "__main__":
    main()
